from docx import Document
from docx.shared import Pt, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from datetime import date
import os

MARGEM_SUPERIOR = Cm(3)
MARGEM_INFERIOR = Cm(2)
MARGEM_ESQUERDA = Cm(3)
MARGEM_DIREITA = Cm(2)
FONTE_NOME = "Times New Roman"
FONTE_TAMANHO = Pt(12)
ESPACAMENTO_LINHA = 1.5

AVISO_LEGAL = (
    "[AVISO] Este documento foi gerado por inteligência artificial com base nos "
    "documentos fornecidos. A revisão por um advogado responsável é obrigatória "
    "antes de qualquer utilização."
)


def criar_documento(titulo: str) -> Document:
    doc = Document()

    for section in doc.sections:
        section.top_margin = MARGEM_SUPERIOR
        section.bottom_margin = MARGEM_INFERIOR
        section.left_margin = MARGEM_ESQUERDA
        section.right_margin = MARGEM_DIREITA

    estilo_normal = doc.styles["Normal"]
    estilo_normal.font.name = FONTE_NOME
    estilo_normal.font.size = FONTE_TAMANHO
    estilo_normal.paragraph_format.line_spacing = ESPACAMENTO_LINHA
    estilo_normal.paragraph_format.space_after = Pt(0)

    return doc


def adicionar_titulo(doc: Document, texto: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(texto)
    run.bold = True
    run.font.size = Pt(14)
    run.font.name = FONTE_NOME
    p.paragraph_format.space_after = Pt(24)
    p.paragraph_format.line_spacing = ESPACAMENTO_LINHA


def adicionar_enderecamento(doc: Document, texto: str = None):
    if texto is None:
        texto = "EXCELENTÍSSIMO SENHOR DOUTOR JUIZ DE DIREITO DA ___ VARA CÍVEL DA COMARCA DE ___"
    p = doc.add_paragraph()
    run = p.add_run(texto)
    run.bold = True
    run.font.size = FONTE_TAMANHO
    run.font.name = FONTE_NOME
    p.paragraph_format.space_after = Pt(18)
    p.paragraph_format.line_spacing = ESPACAMENTO_LINHA


def adicionar_qualificacao(doc: Document, label: str, texto: str):
    if not texto or texto == "Não informado":
        return
    p = doc.add_paragraph()
    run_label = p.add_run(f"{label}: ")
    run_label.bold = True
    run_label.font.size = FONTE_TAMANHO
    run_label.font.name = FONTE_NOME
    run_texto = p.add_run(texto)
    run_texto.font.size = FONTE_TAMANHO
    run_texto.font.name = FONTE_NOME
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = ESPACAMENTO_LINHA


def adicionar_secao(doc: Document, titulo_secao: str, texto: str):
    if not texto:
        return
    p = doc.add_paragraph()
    run = p.add_run(titulo_secao)
    run.bold = True
    run.font.size = FONTE_TAMANHO
    run.font.name = FONTE_NOME
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = ESPACAMENTO_LINHA

    if texto:
        p2 = doc.add_paragraph()
        run2 = p2.add_run(texto)
        run2.font.size = FONTE_TAMANHO
        run2.font.name = FONTE_NOME
        p2.paragraph_format.space_after = Pt(6)
        p2.paragraph_format.line_spacing = ESPACAMENTO_LINHA
        p2.paragraph_format.first_line_indent = Cm(2)


def adicionar_lista(doc: Document, titulo: str, itens: list):
    if not itens:
        return
    p = doc.add_paragraph()
    run = p.add_run(titulo)
    run.bold = True
    run.font.size = FONTE_TAMANHO
    run.font.name = FONTE_NOME
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = ESPACAMENTO_LINHA

    for item in itens:
        p2 = doc.add_paragraph(style="List Bullet")
        run2 = p2.add_run(str(item))
        run2.font.size = FONTE_TAMANHO
        run2.font.name = FONTE_NOME
        p2.paragraph_format.space_after = Pt(3)
        p2.paragraph_format.line_spacing = ESPACAMENTO_LINHA


def adicionar_paragrafo(doc: Document, texto: str, bold: bool = False, alinhar_direita: bool = False):
    if not texto:
        return
    p = doc.add_paragraph()
    run = p.add_run(texto)
    run.bold = bold
    run.font.size = FONTE_TAMANHO
    run.font.name = FONTE_NOME
    if alinhar_direita:
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = ESPACAMENTO_LINHA
    if not bold and not alinhar_direita:
        p.paragraph_format.first_line_indent = Cm(2)


def adicionar_aviso_legal(doc: Document):
    p = doc.add_paragraph()
    run = p.add_run(f"\n{AVISO_LEGAL}")
    run.italic = True
    run.font.size = Pt(9)
    run.font.name = FONTE_NOME
    run.font.color.rgb = None
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)


def adicionar_local_data(doc: Document, cidade: str = None):
    if cidade is None:
        cidade = "[CIDADE]"
    hoje = date.today().strftime("%d de %B de %Y")
    p = doc.add_paragraph()
    run = p.add_run(f"{cidade}, {hoje}.")
    run.font.size = FONTE_TAMANHO
    run.font.name = FONTE_NOME
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = ESPACAMENTO_LINHA


def adicionar_advogado(doc: Document, nome: str = None):
    if nome is None:
        nome = "[NOME DO ADVOGADO]"
    oab = "[NÚMERO OAB]"
    p = doc.add_paragraph()
    run = p.add_run(f"\n{nome}\nOAB {oab}")
    run.font.size = FONTE_TAMANHO
    run.font.name = FONTE_NOME
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.line_spacing = ESPACAMENTO_LINHA


def adicionar_rodape(doc: Document):
    for section in doc.sections:
        footer = section.footer
        footer.is_linked_to_previous = False
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("Nex-Júris - Documento gerado por IA | Revisão obrigatória pelo advogado")
        run.font.size = Pt(8)
        run.font.name = FONTE_NOME
        run.font.color.rgb = None
        run.italic = True


def salvar_docx(doc: Document, filename: str) -> str:
    os.makedirs(os.path.dirname(filename), exist_ok=True)
    doc.save(filename)
    return filename
